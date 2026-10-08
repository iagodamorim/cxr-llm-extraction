"""
merge_v3_groundtruth.py
=========================

Final step of the ground-truth construction. Consumes:
  - reconstructed v1 GT (output of reconstruct_ground_truth_v1.py)
  - reviewed consolidation/pneumonia split (output of re_annotate_v3_helper.py
    after the user has filled in the USER_consolidation / USER_pneumonia
    columns in REVIEW_NEEDED sheet)

Produces ground_truth_v3.csv with the final 8-class schema:
  accession_number, split, opacity, consolidation, pneumonia,
  cardiomegaly, effusion, pneumothorax, mass, yesfinding,
  abnormalities_found

Logic:
  - Take all rows from v1 GT
  - For each row:
      * If accession_number is in REVIEW_NEEDED sheet:
          consolidation = USER_consolidation
          pneumonia     = USER_pneumonia
      * If accession_number is in AUTO_BOTH_ZERO sheet:
          consolidation = 0
          pneumonia     = 0
  - Drop the v1 'bacterial_pneumonia' column (now replaced by consolidation + pneumonia)
  - Recompute yesfinding to ensure it includes the new consolidation/pneumonia signals:
      yesfinding = 1 if ANY of (opacity, consolidation, pneumonia, cardiomegaly,
                                effusion, pneumothorax, mass) == 1
                     OR if the original v1 yesfinding was 1 (which captures
                     unlisted intrathoracic findings)

Author: I.P. D'Amorim et al. PAP v2.1.
"""

import argparse
import sys
from pathlib import Path
import pandas as pd
import openpyxl

V3_NON_COMPOSITE = ["opacity", "consolidation", "pneumonia", "cardiomegaly",
                    "effusion", "pneumothorax", "mass"]


def read_review_xlsx(path: Path) -> tuple[dict, set]:
    """Return ({accession: (cons, pneu, notes)}, set_of_auto_both_zero_accessions)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    review = {}
    if "REVIEW_NEEDED" in wb.sheetnames:
        sh = wb["REVIEW_NEEDED"]
        hdr = [c.value for c in next(sh.iter_rows(min_row=1, max_row=1))]
        idx = {h: i for i, h in enumerate(hdr)}
        for row in sh.iter_rows(min_row=2, values_only=True):
            acc = row[idx["accession_number"]]
            if acc is None: continue
            cons = row[idx["USER_consolidation"]]
            pneu = row[idx["USER_pneumonia"]]
            notes = row[idx.get("USER_notes", -1)] if "USER_notes" in idx else ""
            try:
                cons = int(cons) if cons is not None else 0
                pneu = int(pneu) if pneu is not None else 0
            except (ValueError, TypeError):
                cons, pneu = 0, 0
            review[str(acc)] = (cons, pneu, notes or "")
    auto = set()
    if "AUTO_BOTH_ZERO" in wb.sheetnames:
        sh = wb["AUTO_BOTH_ZERO"]
        for row in sh.iter_rows(min_row=2, values_only=True):
            if row[0] is not None:
                auto.add(str(row[0]))
    return review, auto


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--v1-gt-csv", required=True,
                    help="output of reconstruct_ground_truth_v1.py")
    ap.add_argument("--review-xlsx", required=True,
                    help="filled review xlsx from re_annotate_v3_helper.py")
    ap.add_argument("--output", default="ground_truth_v3.csv")
    args = ap.parse_args()

    v1 = pd.read_csv(args.v1_gt_csv)
    review, auto = read_review_xlsx(Path(args.review_xlsx))

    print(f"v1 GT rows: {len(v1)}")
    print(f"REVIEW_NEEDED rows: {len(review)}")
    print(f"AUTO_BOTH_ZERO rows: {len(auto)}")

    rows = []
    missing = []
    for _, r in v1.iterrows():
        acc = str(r["accession_number"])
        if acc in review:
            cons, pneu, notes = review[acc]
        elif acc in auto:
            cons, pneu, notes = 0, 0, ""
        else:
            missing.append(acc)
            cons, pneu, notes = 0, 0, ""

        # Compute yesfinding fresh
        binary = {c: int(r.get(c, 0)) for c in ["opacity", "cardiomegaly",
                                                "effusion", "pneumothorax", "mass"]}
        binary["consolidation"] = cons
        binary["pneumonia"] = pneu
        yes_from_classes = max(binary.values()) if binary else 0
        # Preserve "yesfinding via unlisted abnormality" from v1
        v1_yes = int(r.get("yesfinding", 0))
        yesfinding = 1 if (yes_from_classes == 1 or v1_yes == 1) else 0

        rows.append({
            "accession_number": acc,
            "split": r["split"],
            "opacity":       binary["opacity"],
            "consolidation": binary["consolidation"],
            "pneumonia":     binary["pneumonia"],
            "cardiomegaly":  binary["cardiomegaly"],
            "effusion":      binary["effusion"],
            "pneumothorax":  binary["pneumothorax"],
            "mass":          binary["mass"],
            "yesfinding":    yesfinding,
            "abnormalities_found": r.get("abnormalities_found", ""),
            "review_notes":  notes,
        })

    if missing:
        print(f"WARN: {len(missing)} accessions present in v1 GT but missing from review xlsx; assigned cons=0, pneu=0.")

    out = pd.DataFrame(rows)
    out.to_csv(args.output, index=False)

    n = len(out)
    print(f"\nWrote: {args.output} ({n} rows)")
    print("\nFinal v3 class prevalence:")
    for c in V3_NON_COMPOSITE + ["yesfinding"]:
        pos = int(out[c].sum())
        print(f"  {c:>15s}  {pos:>5d} / {n}  ({100*pos/n:5.2f}%)")
    print("\nSplit summary:")
    for s in ["dev", "test"]:
        print(f"  {s}: {int((out['split']==s).sum())}")


if __name__ == "__main__":
    main()
