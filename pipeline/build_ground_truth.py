"""
build_ground_truth.py
======================

Builds the 7-class ground truth from the two human-validated xlsx files
(dev + test). One row per accession.

Per-class rule:
    GT[c] = (LLM_extracted[c] == 1 AND c NOT IN FP) OR (c IN FN)

Inputs handle both formats transparently:
  - pipe   : "opacity:0 | cardiomegaly:1 | ... | yesfinding:1 (free text)"
  - JSON   : {"exam_type":..., "findings":[{"presence":"0","class_name":"opacity"}, ...]}

The 7-class slot `pneumonia` absorbs any row whose extraction
called the class `consolidation` (mid-run schema drift in test_780).
The original label is preserved in `pneumonia_label_original` for the
later 8-class split.

Outputs:
  ground_truth_7class.csv  -- one row per accession with the GT
  ground_truth_audit.xlsx  -- only rows that need a human eyeball
"""

import argparse
import json
import re
from pathlib import Path
from collections import Counter

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

V1_CLASSES = ["opacity", "cardiomegaly", "effusion", "mass", "pneumothorax",
              "pneumonia", "yesfinding"]

# typos / aliases we silently normalize to the v4 canonical name
ALIASES = {
    "cadiomegaly": "cardiomegaly",
    "consolidation": "pneumonia",       # v1 schema drift → v4 pneumonia
    "bacterial_pneumonia": "pneumonia", # v1/v3 legacy name → v4 broader pneumonia
}


def clean(s):
    if s is None: return ""
    return str(s).replace("_x000d_", "").strip()


def parse_class_set(s):
    """Parse FP/FN strings like 'opacity;\\nyesfinding;' into (normalized_set, raw_tokens)."""
    if not s: return set(), []
    out, raw = set(), []
    for tok in re.split(r"[;,\n]", clean(s)):
        tok = tok.strip().lower()
        if not tok: continue
        raw.append(tok)
        tok = ALIASES.get(tok, tok)
        if tok in V1_CLASSES:
            out.add(tok)
    return out, raw


def parse_extracted(s):
    """Auto-detect pipe vs JSON. Returns (per_class_dict, abnormalities_found, pneumonia_label_original)."""
    if not s: return {}, "", ""
    s = clean(s)
    if s.startswith("{") or s.startswith("["):
        return _parse_json(s)
    return _parse_pipe(s)


def _parse_json(s):
    try:
        d = json.loads(s)
    except Exception:
        return {}, "", ""
    out, abn, pneu_label = {}, "", ""
    for f in d.get("findings", []):
        if not isinstance(f, dict): continue
        cname = str(f.get("class_name", "")).strip().lower()
        try:
            presence = int(str(f.get("presence", "0")))
        except (ValueError, TypeError):
            presence = 0
        if cname in ("pneumonia", "consolidation", "bacterial_pneumonia"):
            if presence == 1 and not pneu_label:
                pneu_label = cname
            cname = "pneumonia"
        out[cname] = presence
        if cname == "yesfinding":
            abn = f.get("abnormalities_found", "") or ""
    return out, abn, pneu_label


def _parse_pipe(s):
    out, abn, pneu_label = {}, "", ""
    for part in s.split("|"):
        part = part.strip()
        if ":" not in part: continue
        k, rest = part.split(":", 1)
        k = k.strip().lower()
        m = re.match(r"^([01])(?:\s*\((.*)\))?\s*$", rest.strip(), flags=re.DOTALL)
        if not m: continue
        presence = int(m.group(1))
        if k in ("pneumonia", "consolidation", "bacterial_pneumonia"):
            if presence == 1 and not pneu_label:
                pneu_label = k
            k = "pneumonia"
        out[k] = presence
        if k == "yesfinding" and m.group(2):
            abn = m.group(2).strip()
    return out, abn, pneu_label


def reconstruct(ext, fp, fn):
    gt = {}
    for c in V1_CLASSES:
        e = ext.get(c, 0)
        if c in fn:    gt[c] = 1
        elif c in fp:  gt[c] = 0
        else:          gt[c] = e
    return gt


def load_xlsx(path, split):
    """Yield dicts: accession, report, extracted, TRUE, FP, FN, Observações."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    hdrs = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    # tolerate older column names
    aliases = {
        "ground_truth": "extracted",
        "VP": "TRUE",
        "FP = classes que foram extraídas e não existem": "FP",
        "FN = classes que existem mas não foram extraídas.": "FN",
    }
    hdrs = [aliases.get(h, h) for h in hdrs]
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None: continue
        acc = str(row[0]).strip()
        if not acc or acc in seen: continue
        seen.add(acc)
        d = dict(zip(hdrs, row))
        d["accession_number"] = acc
        d["split"] = split
        d["source_file"] = Path(path).name
        yield d


def process(path, split):
    out_rows, audit_rows = [], []
    counts = Counter()
    for d in load_xlsx(path, split):
        ext, abn, pneu_label = parse_extracted(d.get("extracted"))
        fp, fp_raw = parse_class_set(d.get("FP"))
        fn, fn_raw = parse_class_set(d.get("FN"))
        gt = reconstruct(ext, fp, fn)

        true_str = clean(d.get("TRUE"))
        is_ok = (true_str.upper() == "OK") and not fp and not fn
        kind = "OK" if is_ok else "corrected"
        counts[kind] += 1

        # ---- audit conditions ---------------------------------------
        notes = []
        # tokens user typed that we couldn't match to any class
        unknown_fp = [t for t in fp_raw if ALIASES.get(t, t) not in V1_CLASSES]
        unknown_fn = [t for t in fn_raw if ALIASES.get(t, t) not in V1_CLASSES]
        if unknown_fp: notes.append(f"FP token desconhecido: {unknown_fp}")
        if unknown_fn: notes.append(f"FN token desconhecido: {unknown_fn}")
        # FP refers to a class that wasn't extracted=1
        for c in fp:
            if ext.get(c, 0) != 1:
                notes.append(f"FP={c} mas extração era 0 (contradição)")
        # FN refers to a class that was already extracted=1
        for c in fn:
            if ext.get(c, 0) == 1:
                notes.append(f"FN={c} mas extração já era 1 (redundante)")
        # TRUE list doesn't match (extracted - FP)
        if not is_ok and true_str:
            true_set, _ = parse_class_set(true_str)
            derived = {c for c in V1_CLASSES if ext.get(c, 0) == 1 and c not in fp}
            if true_set and true_set != derived:
                notes.append(f"TRUE={sorted(true_set)} != (extraído-FP)={sorted(derived)}")
        # yesfinding=1 sem nada
        any_bin = any(gt[c] == 1 for c in V1_CLASSES if c != "yesfinding")
        if gt["yesfinding"] == 1 and not any_bin and not abn:
            notes.append("yesfinding=1 sem classe binária positiva e sem abnormalities_found")

        out_rows.append({
            "accession_number": d["accession_number"],
            "split": split,
            "report": d.get("report") or "",
            **gt,
            "abnormalities_found": abn,
            "pneumonia_label_original": pneu_label,
            "source_file": d["source_file"],
            "n_fp": len(fp),
            "n_fn": len(fn),
            "validation_kind": kind,
        })

        if notes:
            audit_rows.append({
                "accession_number": d["accession_number"],
                "split": split,
                "validation_kind": kind,
                "issues": " | ".join(notes),
                "extracted_raw": clean(d.get("extracted"))[:300],
                "TRUE": true_str,
                "FP": clean(d.get("FP")),
                "FN": clean(d.get("FN")),
                "Observações": clean(d.get("Observações")),
            })

    return out_rows, audit_rows, counts


def write_csv(path, rows):
    import csv
    cols = ["accession_number", "split", "report",
            *V1_CLASSES, "abnormalities_found",
            "pneumonia_label_original", "source_file",
            "n_fp", "n_fn", "validation_kind"]
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows: w.writerow(r)


def write_audit_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "AUDIT"
    cols = ["accession_number", "split", "validation_kind", "issues",
            "extracted_raw", "TRUE", "FP", "FN", "Observações"]
    for ci, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=ci, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1F3864")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for ri, r in enumerate(rows, 2):
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(row=ri, column=ci, value=r.get(c, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {"accession_number": 16, "split": 8, "validation_kind": 12,
              "issues": 60, "extracted_raw": 50, "TRUE": 30, "FP": 25,
              "FN": 25, "Observações": 30}
    from openpyxl.utils import get_column_letter
    for ci, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(ci)].width = widths.get(c, 18)
    ws.freeze_panes = "A2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dev-xlsx",
                    default="data/uploads/dev_295_validated_v2.xlsx")
    ap.add_argument("--test-xlsx",
                    default="data/uploads/test_780_validated.xlsx")
    ap.add_argument("--output-csv",
                    default="data/outputs/ground_truth_7class.csv")
    ap.add_argument("--audit-xlsx",
                    default="data/outputs/ground_truth_audit.xlsx")
    args = ap.parse_args()

    dev_rows, dev_audit, dev_counts = process(args.dev_xlsx, "dev")
    test_rows, test_audit, test_counts = process(args.test_xlsx, "test")

    all_rows = dev_rows + test_rows
    all_audit = dev_audit + test_audit

    Path(args.output_csv).parent.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_csv, all_rows)
    write_audit_xlsx(args.audit_xlsx, all_audit)

    print(f"dev  : {len(dev_rows)} rows   OK={dev_counts['OK']}  corrigidos={dev_counts['corrected']}")
    print(f"test : {len(test_rows)} rows  OK={test_counts['OK']}  corrigidos={test_counts['corrected']}")
    print(f"total: {len(all_rows)} rows")
    print(f"\nWrote: {args.output_csv}")
    print(f"Wrote: {args.audit_xlsx}  ({len(all_audit)} linhas pra revisar)")

    print("\nPrevalência por classe (no ground-truth):")
    n = len(all_rows)
    for c in V1_CLASSES:
        pos = sum(1 for r in all_rows if r[c] == 1)
        print(f"  {c:>22s}  {pos:>5d} / {n}  ({100*pos/n:5.2f}%)")

    pneu_origins = Counter(r["pneumonia_label_original"]
                            for r in all_rows
                            if r["pneumonia"] == 1)
    print(f"\nDentro de pneumonia=1 ({sum(pneu_origins.values())} casos):")
    for k, v in pneu_origins.most_common():
        print(f"  rotulado pelo LLM como {k or '(vazio)'}: {v}")


if __name__ == "__main__":
    main()
