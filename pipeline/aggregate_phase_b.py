"""
aggregate_phase_b.py
=====================

Phase B gate: aggregate per-vendor metrics (output of compare_predictions.py)
and decide the Specialist Routing assignment from the dev split.

This script is READ-ONLY on existing artifacts. It does not call APIs,
does not modify any response or metrics CSV, and produces three new
artifacts under data/outputs/:

  - phase_b_summary.csv         (long-form: vendor × class × metric)
  - phase_b_summary.xlsx        (pretty matrix + routing + go/no-go)
  - specialist_routing_assignment.json   (PAP-compliant routing decision)

Pipeline:
  1. Discover metrics CSVs matching --metrics-pattern (default: metrics_*_dev.csv).
     Smoke-test files (containing "smoke") are skipped.
  2. Extract vendor name from filename: metrics_{VENDOR}_dev.csv -> VENDOR.
  3. Build a wide matrix: vendor (rows) × class (cols) of dev-split F1.
  4. Compute macro F1 per vendor.
  5. For each class, identify the vendor with the highest F1 -> Specialist
     Routing assignment. Stable-sort tie-break (alphabetical vendor name).
  6. (Optional) read response CSVs from --responses-dir to compute
     parse_ok rate and entailment_violation rate per vendor (quality metrics).
  7. Run go/no-go check per PAP v2.1 criteria:
        - Per-vendor macro F1 >= 0.80
        - Per-vendor pneumonia F1 >= 0.65
        - Per-vendor parse_ok >= 0.98 (if responses available)
  8. Print a formatted report to stdout.

Usage:
  python pipeline/aggregate_phase_b.py
  python pipeline/aggregate_phase_b.py --metrics-dir data/outputs \\
                                       --responses-dir responses

Outputs to:
  data/outputs/phase_b_summary.csv
  data/outputs/phase_b_summary.xlsx
  data/outputs/specialist_routing_assignment.json

Author: IPDA et al. PAP v2.1, Phase B gate (May 2026).
"""

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from glob import glob
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

V4_CLASSES = ["opacity", "pneumonia", "cardiomegaly", "effusion",
              "pneumothorax", "mass", "yesfinding"]

# PAP v2.1 go/no-go thresholds
GATE_MACRO_F1 = 0.80
GATE_PNEUMONIA_F1 = 0.65
GATE_PARSE_OK = 0.98


# ----- File discovery -------------------------------------------------------

def discover_metrics_files(metrics_dir: Path, pattern: str) -> dict[str, Path]:
    """Return {vendor: path}, skipping smoke-test files."""
    files = sorted(glob(str(metrics_dir / pattern)))
    out = {}
    for f in files:
        name = Path(f).name
        if "smoke" in name.lower():
            continue
        # metrics_{VENDOR}_{split}.csv  ->  VENDOR  (split = dev OR test)
        m = re.match(r"metrics_(.+)_(?:dev|test)\.csv$", name)
        if not m:
            continue
        vendor = m.group(1)
        out[vendor] = Path(f)
    return out


def discover_response_files(responses_dir: Path) -> dict[str, Path]:
    """Return {vendor: path} for response CSVs, skipping smoke-test files."""
    if not responses_dir.exists():
        return {}
    files = sorted(glob(str(responses_dir / "*_dev.csv")))
    out = {}
    for f in files:
        name = Path(f).name
        if "smoke" in name.lower():
            continue
        m = re.match(r"(.+)_dev\.csv$", name)
        if not m:
            continue
        out[m.group(1)] = Path(f)
    return out


# ----- Metrics ingestion ----------------------------------------------------

def load_vendor_metrics(path: Path, split: str = "dev") -> pd.DataFrame:
    """Read one metrics_{vendor}_{split}.csv and return only that split's rows."""
    df = pd.read_csv(path)
    if "split" not in df.columns or "classe" not in df.columns or "f1" not in df.columns:
        raise ValueError(f"{path}: missing expected columns (split/classe/f1)")
    return df[df["split"] == split].copy()


def build_f1_matrix(vendor_metrics: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Build vendor × class wide DataFrame of dev-split F1 scores."""
    rows = {}
    for vendor, df in vendor_metrics.items():
        rows[vendor] = {row["classe"]: row["f1"] for _, row in df.iterrows()}
    mat = pd.DataFrame(rows).T
    # Ensure consistent class order
    cols_present = [c for c in V4_CLASSES if c in mat.columns]
    return mat[cols_present]


def compute_macro_f1(f1_matrix: pd.DataFrame) -> pd.Series:
    return f1_matrix.mean(axis=1).round(4)


# ----- Specialist Routing decision ------------------------------------------

def decide_routing(f1_matrix: pd.DataFrame) -> dict:
    """For each class, return the vendor with the highest F1.

    Tie-break: alphabetical vendor name (stable, reproducible).
    """
    assignment = {}
    for cls in f1_matrix.columns:
        col = f1_matrix[cls]
        # Sort by (-F1, vendor) to get stable max with alphabetical tie-break
        sorted_vendors = sorted(col.index, key=lambda v: (-col[v], v))
        best_vendor = sorted_vendors[0]
        # Find any ties for transparency
        best_f1 = col[best_vendor]
        ties = [v for v in col.index if abs(col[v] - best_f1) < 1e-9]
        assignment[cls] = {
            "best_vendor": best_vendor,
            "best_f1": round(float(best_f1), 4),
            "ties": ties if len(ties) > 1 else [],
            "ranking": [(v, round(float(col[v]), 4)) for v in sorted_vendors],
        }
    return assignment


# ----- Quality metrics (from response CSVs) ---------------------------------

def vendor_quality(response_path: Path) -> dict:
    """Return parse_ok and entailment_violation rates from a response CSV."""
    df = pd.read_csv(response_path)
    n = len(df)
    if n == 0:
        return {"n": 0, "parse_ok_rate": None, "entailment_violation_rate": None}
    parse_ok = df["parse_ok"].astype(bool).mean() if "parse_ok" in df.columns else None
    if "entailment_violation" in df.columns:
        ev = df["entailment_violation"].astype(bool).mean()
    else:
        ev = None
    return {
        "n": int(n),
        "parse_ok_rate": round(float(parse_ok), 4) if parse_ok is not None else None,
        "entailment_violation_rate": round(float(ev), 4) if ev is not None else None,
    }


# ----- Go/no-go check -------------------------------------------------------

def go_no_go(f1_matrix: pd.DataFrame, macro_f1: pd.Series,
             quality: dict[str, dict]) -> dict:
    """Apply PAP v2.1 gate criteria. Return per-vendor checks + overall status."""
    checks = {}
    for vendor in f1_matrix.index:
        c = {}
        c["macro_f1"] = round(float(macro_f1[vendor]), 4)
        c["macro_f1_pass"] = bool(c["macro_f1"] >= GATE_MACRO_F1)
        pneu_f1 = f1_matrix.loc[vendor, "pneumonia"] if "pneumonia" in f1_matrix.columns else None
        c["pneumonia_f1"] = round(float(pneu_f1), 4) if pneu_f1 is not None else None
        c["pneumonia_pass"] = bool(pneu_f1 is not None and pneu_f1 >= GATE_PNEUMONIA_F1)
        q = quality.get(vendor, {})
        c["parse_ok"] = q.get("parse_ok_rate")
        c["parse_ok_pass"] = bool((c["parse_ok"] is None) or (c["parse_ok"] >= GATE_PARSE_OK))
        c["all_pass"] = bool(c["macro_f1_pass"] and c["pneumonia_pass"] and c["parse_ok_pass"])
        checks[vendor] = c
    overall = bool(all(c["all_pass"] for c in checks.values())) if checks else False
    return {"per_vendor": checks, "overall_go": overall}


# ----- Output writers -------------------------------------------------------

def write_summary_csv(f1_matrix: pd.DataFrame, macro_f1: pd.Series,
                       quality: dict[str, dict], out_path: Path) -> None:
    rows = []
    for vendor in f1_matrix.index:
        for cls in f1_matrix.columns:
            rows.append({
                "vendor": vendor,
                "metric": "f1",
                "classe": cls,
                "value": round(float(f1_matrix.loc[vendor, cls]), 4),
            })
        rows.append({"vendor": vendor, "metric": "macro_f1",
                     "classe": "—", "value": round(float(macro_f1[vendor]), 4)})
        q = quality.get(vendor, {})
        rows.append({"vendor": vendor, "metric": "parse_ok_rate",
                     "classe": "—", "value": q.get("parse_ok_rate")})
        rows.append({"vendor": vendor, "metric": "entailment_violation_rate",
                     "classe": "—", "value": q.get("entailment_violation_rate")})
    pd.DataFrame(rows).to_csv(out_path, index=False)


def write_summary_xlsx(f1_matrix: pd.DataFrame, macro_f1: pd.Series,
                        quality: dict[str, dict], routing: dict,
                        gate: dict, out_path: Path) -> None:
    wb = Workbook()
    NAVY = "1F3864"; LIGHT = "F2F2F2"; GOLD = "FFE699"; GREEN = "C6EFCE"
    AMBER = "FFEB9C"; RED = "FFC7CE"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # --- Sheet 1: F1 matrix with best-cell highlight ---
    ws = wb.active
    ws.title = "F1 matrix"
    hdr = ["vendor", *list(f1_matrix.columns), "macro_F1"]
    for ci, h in enumerate(hdr, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    ws.row_dimensions[1].height = 22

    # Identify best vendor per class
    best_per_class = {cls: routing[cls]["best_vendor"] for cls in f1_matrix.columns}

    for ri, vendor in enumerate(f1_matrix.index, 2):
        ws.cell(row=ri, column=1, value=vendor).font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=ri, column=1).border = border
        for ci, cls in enumerate(f1_matrix.columns, 2):
            v = float(f1_matrix.loc[vendor, cls])
            c = ws.cell(row=ri, column=ci, value=round(v, 4))
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(horizontal="center")
            c.border = border
            if best_per_class.get(cls) == vendor:
                c.fill = PatternFill("solid", start_color=GOLD)
                c.font = Font(name="Arial", size=10, bold=True)
        mc = ws.cell(row=ri, column=len(hdr), value=round(float(macro_f1[vendor]), 4))
        mc.font = Font(name="Arial", size=10, bold=True)
        mc.alignment = Alignment(horizontal="center")
        mc.border = border
        if ri % 2 == 0:
            for ci in range(1, len(hdr) + 1):
                if not ws.cell(row=ri, column=ci).fill.start_color.rgb or \
                   ws.cell(row=ri, column=ci).fill.start_color.rgb in ("00000000",):
                    ws.cell(row=ri, column=ci).fill = PatternFill("solid", start_color=LIGHT)

    for ci, w in enumerate([18, 12, 12, 12, 12, 12, 12, 12, 12], 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "B2"

    # --- Sheet 2: Specialist Routing assignment ---
    ws2 = wb.create_sheet("Specialist Routing")
    hdr2 = ["classe", "best_vendor", "best_f1", "ties", "full_ranking"]
    for ci, h in enumerate(hdr2, 1):
        c = ws2.cell(row=1, column=ci, value=h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for ri, cls in enumerate(f1_matrix.columns, 2):
        info = routing[cls]
        ws2.cell(row=ri, column=1, value=cls)
        ws2.cell(row=ri, column=2, value=info["best_vendor"]).fill = PatternFill("solid", start_color=GOLD)
        ws2.cell(row=ri, column=3, value=info["best_f1"])
        ws2.cell(row=ri, column=4, value=", ".join(info["ties"]) if info["ties"] else "")
        ws2.cell(row=ri, column=5, value=" | ".join(f"{v}={f1}" for v, f1 in info["ranking"]))
        for ci in range(1, len(hdr2) + 1):
            ws2.cell(row=ri, column=ci).font = Font(name="Arial", size=10)
            ws2.cell(row=ri, column=ci).border = border
    for ci, w in enumerate([16, 18, 10, 24, 70], 1):
        ws2.column_dimensions[get_column_letter(ci)].width = w
    ws2.freeze_panes = "A2"

    # --- Sheet 3: Quality + Go/no-go ---
    ws3 = wb.create_sheet("Gate (go-no-go)")
    hdr3 = ["vendor", "macro_F1", "macro_F1_pass", "pneumonia_F1", "pneumonia_pass",
            "parse_ok_rate", "parse_ok_pass", "entailment_violation_rate", "ALL_PASS"]
    for ci, h in enumerate(hdr3, 1):
        c = ws3.cell(row=1, column=ci, value=h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws3.row_dimensions[1].height = 36
    for ri, (vendor, chk) in enumerate(gate["per_vendor"].items(), 2):
        q = quality.get(vendor, {})
        ws3.cell(row=ri, column=1, value=vendor).font = Font(name="Arial", size=10, bold=True)
        ws3.cell(row=ri, column=2, value=chk["macro_f1"])
        ws3.cell(row=ri, column=3, value="PASS" if chk["macro_f1_pass"] else "FAIL")
        ws3.cell(row=ri, column=4, value=chk["pneumonia_f1"])
        ws3.cell(row=ri, column=5, value="PASS" if chk["pneumonia_pass"] else "FAIL")
        ws3.cell(row=ri, column=6, value=chk["parse_ok"])
        ws3.cell(row=ri, column=7, value="PASS" if chk["parse_ok_pass"] else
                                          ("N/A" if chk["parse_ok"] is None else "FAIL"))
        ws3.cell(row=ri, column=8, value=q.get("entailment_violation_rate"))
        all_pass_cell = ws3.cell(row=ri, column=9, value="GO" if chk["all_pass"] else "NO-GO")
        all_pass_cell.font = Font(name="Arial", size=10, bold=True)
        all_pass_cell.fill = PatternFill("solid", start_color=GREEN if chk["all_pass"] else AMBER)
        for ci in range(1, len(hdr3) + 1):
            cell = ws3.cell(row=ri, column=ci)
            if not cell.font.bold:
                cell.font = Font(name="Arial", size=10)
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
    for ci, w in enumerate([18, 10, 12, 12, 14, 12, 12, 22, 11], 1):
        ws3.column_dimensions[get_column_letter(ci)].width = w
    ws3.freeze_panes = "B2"

    wb.save(out_path)


def write_routing_json(routing: dict, gate: dict, vendors: list[str],
                        out_path: Path, split: str = "dev") -> None:
    split_label = {
        "dev": "dev split, N=295 (ECR cohort) — PAP v2.1 compliant",
        "test": "test split, N=780 — WARNING: routing decision on test "
                "violates PAP v2.1 (oracle/overfit). Use only as exploratory comparison.",
        "all": "combined split, N=1075 (dev + test) — exploratory only, "
               "routing decision must use dev alone per PAP v2.1.",
    }.get(split, f"split={split}")
    payload = {
        "schema": "v4 (7-class)",
        "decision_basis": split_label,
        "split": split,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "pap_version": "v2.1",
        "tie_break": "alphabetical vendor name (stable)",
        "vendors_considered": vendors,
        "assignment": {cls: info["best_vendor"] for cls, info in routing.items()},
        "assignment_detail": routing,
        "gate": gate,
        "thresholds": {
            "macro_f1_min": GATE_MACRO_F1,
            "pneumonia_f1_min": GATE_PNEUMONIA_F1,
            "parse_ok_min": GATE_PARSE_OK,
        },
        "locked": True,
        "note": ("This assignment, once locked, must be used as-is for Phase C "
                 "(test split) per PAP v2.1. Any change is a protocol amendment."),
    }
    out_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False))


# ----- Console report -------------------------------------------------------

def print_report(f1_matrix: pd.DataFrame, macro_f1: pd.Series,
                  quality: dict[str, dict], routing: dict, gate: dict,
                  missing_vendors: list[str], split: str = "dev") -> None:
    split_n = {"dev": 295, "test": 780, "all": 1075}.get(split, "?")
    print("\n" + "=" * 78)
    print(f"AGGREGATE METRICS  —  {split} split (N={split_n}), schema v4 (7 classes)")
    print("=" * 78)

    print(f"\nVendors with metrics found ({len(f1_matrix)}):")
    for v in f1_matrix.index:
        print(f"  - {v}")
    if missing_vendors:
        print(f"\nNOTE: vendors expected but no metrics file found:")
        for v in missing_vendors:
            print(f"  - {v}  (will be excluded from this run)")

    print(f"\n--- Per-class F1 ({split} split, N={split_n}) ---")
    print(f1_matrix.round(4).to_string())

    print("\n--- Macro F1 by vendor ---")
    for v in macro_f1.sort_values(ascending=False).index:
        print(f"  {v:<24s}  {macro_f1[v]:.4f}")

    print("\n--- Quality metrics ---")
    print(f"{'vendor':<24s} {'parse_ok':>10s}   {'entailment_violation':>22s}")
    for v in f1_matrix.index:
        q = quality.get(v, {})
        po = q.get("parse_ok_rate")
        ev = q.get("entailment_violation_rate")
        po_s = f"{po:.4f}" if po is not None else "  N/A   "
        ev_s = f"{ev:.4f}" if ev is not None else "  N/A   "
        print(f"  {v:<22s}  {po_s:>10s}   {ev_s:>22s}")

    print("\n--- Specialist Routing assignment ---")
    print(f"{'classe':<16s}  {'best_vendor':<24s}  {'F1':>6s}  ties")
    for cls in f1_matrix.columns:
        info = routing[cls]
        ties_s = f"(ties: {', '.join(info['ties'])})" if info["ties"] else ""
        print(f"  {cls:<14s}  {info['best_vendor']:<24s}  {info['best_f1']:.4f}  {ties_s}")

    print("\n--- Go/no-go check (PAP v2.1) ---")
    print(f"  Macro F1 threshold: >= {GATE_MACRO_F1}")
    print(f"  Pneumonia F1 threshold: >= {GATE_PNEUMONIA_F1}")
    print(f"  Parse OK threshold: >= {GATE_PARSE_OK}")
    print()
    for v, chk in gate["per_vendor"].items():
        status = "GO " if chk["all_pass"] else "HOLD"
        print(f"  [{status}] {v:<24s}  macro_F1={chk['macro_f1']:.4f}  pneu_F1={chk['pneumonia_f1']}  parse_ok={chk['parse_ok']}")

    print("\n" + "=" * 78)
    if gate["overall_go"]:
        print("OVERALL: GREEN LIGHT — proceed to Phase C (test 780)")
    else:
        failed = [v for v, c in gate["per_vendor"].items() if not c["all_pass"]]
        print(f"OVERALL: HOLD — review {len(failed)} vendor(s) before Phase C:")
        for v in failed:
            print(f"  - {v}")
    print("=" * 78 + "\n")


# ----- Main -----------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--metrics-dir", default="data/outputs")
    ap.add_argument("--metrics-pattern", default="metrics_*_dev.csv")
    ap.add_argument("--split", default=None,
                    help="dev|test|all — auto-derivado do pattern se omitido")
    ap.add_argument("--responses-dir", default="responses")
    ap.add_argument("--output-summary-csv", default="data/outputs/phase_b_summary.csv")
    ap.add_argument("--output-summary-xlsx", default="data/outputs/phase_b_summary.xlsx")
    ap.add_argument("--output-routing", default="data/outputs/specialist_routing_assignment.json")
    ap.add_argument("--expected-vendors", nargs="*",
                    default=["gpt-4-1", "gpt-4-1-mini", "gemini-2-5-pro",
                             "gemini-2-5-flash", "sonnet-4-6", "haiku-4-5"],
                    help="vendor names expected (for missing-files warning)")
    args = ap.parse_args()

    metrics_dir = Path(args.metrics_dir)
    responses_dir = Path(args.responses_dir)

    vendor_files = discover_metrics_files(metrics_dir, args.metrics_pattern)
    if not vendor_files:
        sys.exit(f"ERROR: no metrics files found in {metrics_dir} "
                 f"matching {args.metrics_pattern}")

    missing = sorted(set(args.expected_vendors) - set(vendor_files.keys()))

    # Derive split from pattern if not given
    split = args.split
    if split is None:
        if "_test" in args.metrics_pattern:
            split = "test"
        elif "_dev" in args.metrics_pattern:
            split = "dev"
        else:
            split = "all"
    print(f"Loading split: {split}")

    # Load metrics
    vendor_metrics = {}
    for vendor, path in vendor_files.items():
        try:
            vendor_metrics[vendor] = load_vendor_metrics(path, split=split)
        except Exception as e:
            print(f"WARN: failed to load {path}: {e}", file=sys.stderr)

    f1_matrix = build_f1_matrix(vendor_metrics)
    if f1_matrix.empty:
        sys.exit("ERROR: F1 matrix is empty after loading; abort.")

    macro_f1 = compute_macro_f1(f1_matrix)

    # Load quality (parse_ok, entailment_violation) from response CSVs
    response_files = discover_response_files(responses_dir)
    quality = {}
    for vendor in f1_matrix.index:
        if vendor in response_files:
            try:
                quality[vendor] = vendor_quality(response_files[vendor])
            except Exception as e:
                print(f"WARN: quality read failed for {vendor}: {e}", file=sys.stderr)
                quality[vendor] = {}
        else:
            quality[vendor] = {}

    # Routing decision
    routing = decide_routing(f1_matrix)

    # Go/no-go
    gate = go_no_go(f1_matrix, macro_f1, quality)

    # Write outputs
    Path(args.output_summary_csv).parent.mkdir(parents=True, exist_ok=True)
    write_summary_csv(f1_matrix, macro_f1, quality, Path(args.output_summary_csv))
    write_summary_xlsx(f1_matrix, macro_f1, quality, routing, gate,
                        Path(args.output_summary_xlsx))
    write_routing_json(routing, gate, list(f1_matrix.index),
                       Path(args.output_routing), split=split)

    print_report(f1_matrix, macro_f1, quality, routing, gate, missing, split=split)

    print(f"Wrote: {args.output_summary_csv}")
    print(f"Wrote: {args.output_summary_xlsx}")
    print(f"Wrote: {args.output_routing}")


if __name__ == "__main__":
    main()
